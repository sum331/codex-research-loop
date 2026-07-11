from __future__ import annotations

import json
import os
import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


PLUGIN_ROOT = Path(__file__).resolve().parents[2]
PROJECT_ROOT_ENV = os.environ.get("J1030_PROJECT_ROOT")
if not PROJECT_ROOT_ENV:
    raise SystemExit("Set J1030_PROJECT_ROOT to the J1030 case project root before rebuilding this report.")
PROJECT_ROOT = Path(PROJECT_ROOT_ENV).expanduser().resolve()
OUT_DIR = PLUGIN_ROOT / "docs" / "reports"
ASSET_DIR = OUT_DIR / "assets"
MINDMAP_PNG = ASSET_DIR / "loop_research_mindmap.png"
OUT_MD = OUT_DIR / "codex_research_loop_design_and_j1030_case_report.md"
OUT_DOCX = OUT_DIR / "codex_research_loop_design_and_j1030_case_report_with_mindmap.docx"
OUT_DOCX_LEGACY = OUT_DIR / "codex_research_loop_design_and_j1030_case_report.docx"


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


plugin_manifest = load_json(PLUGIN_ROOT / ".codex-plugin" / "plugin.json")
j1030_state = load_json(PROJECT_ROOT / ".research-loop" / "state.json")
j1030_passport = load_json(PROJECT_ROOT / ".research-loop" / "material-passport.json")
j1030_analysis = load_json(PROJECT_ROOT / "reports" / "j1030_analysis_summary.json")
j1030_stats = load_json(PROJECT_ROOT / "reports" / "j1030_formal_report_stats.json")
j1030_build = load_json(PROJECT_ROOT / "reports" / "j1030_formal_research_report_build_summary.json")
j1030_deep = load_json(
    PROJECT_ROOT / ".research-loop" / "deep-loops" / "20260705-234955-deep-20260705-234955-p8.json"
)


REPORT_DATE = "2026-07-11"


def repo_ref(path: Path) -> str:
    return path.resolve().relative_to(PLUGIN_ROOT).as_posix()


def case_ref(value: str | Path) -> str:
    path = Path(value).expanduser().resolve()
    try:
        rel = path.relative_to(PROJECT_ROOT)
    except ValueError:
        return str(path)
    return "$env:J1030_PROJECT_ROOT\\" + str(rel).replace("/", "\\")


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    out = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        out.append("| " + " | ".join(str(x).replace("\n", "<br>") for x in row) + " |")
    return "\n".join(out)


top5 = j1030_stats["top10"][:5]
top5_rows = [
    [
        f"XID{row['XID']}",
        str(row["candidate_score"]),
        row["candidate_class"],
        f"{row['z']:.4g}",
        row["specType"],
        f"{row['logNH_spectral']:.2f}",
        f"{row['logL210']:.2f}",
    ]
    for row in top5
]

capability_rows = [
    ["状态与护照", ".research-loop/state.json, material-passport.json", "跨线程保存课题、材料、主张、证据、风险、下一步"],
    ["自然语言入口", "normalize", "把低结构输入转化为可执行 prompt、任务类型、链路深度和缺失槽位"],
    ["多通路路由", "route, P1-P10, N0-N9", "按科研阶段选择多个子链路，并让共享节点负责分发和回流"],
    ["存储规范", "storage-policy", "在旧项目中自适应映射，在新项目中初始化 canonical 科研目录"],
    ["素材摄取", "content-ingest, source-hub", "文章进入文章链路，数据进入标准数据包，来源元数据进入证据系统"],
    ["证据阀门", "claim-evidence", "结构化检查 claim 与 evidence 的连接、来源、定位和状态"],
    ["深循环阀门", "deep-loop", "产出 route_next, retry_same_route, escalate_problem_loop, pause_for_human"],
    ["无人值守执行", "auto-loop --auto-route-next", "把 deep-loop 决策交给执行器，触发下一链路或同链路重试"],
    ["问题诊断链", "problem-loop, problem-promote", "在隔离 scratch lab 中生成专家组、测试和调整方案，过闸后才改核心文件"],
    ["外部生态", "Zotero, llm-wiki, Codex CLI, MCP", "支持文献库、知识库、跨线程调用和自动代理启动"],
]

subchain_rows = [
    ["P1", "scoping-question-chain", "问题定义、边界、研究目标"],
    ["P2", "literature-evidence-chain", "文献发现、文章摄取、证据记录"],
    ["P3", "claim-contribution-chain", "主张、贡献、创新性和证据对应"],
    ["P4", "method-experiment-compliance-chain", "方法设计、实验协议、合规检查"],
    ["P5", "data-code-execution-chain", "数据、代码、运行日志和复现"],
    ["P6", "analysis-statistics-figure-chain", "分析、统计、表格和图形"],
    ["P7", "writing-citation-format-chain", "写作、引用、DOCX/LaTeX/PDF"],
    ["P8", "review-revision-integrity-chain", "评审、修订、完整性阀门"],
    ["P9", "submission-publication-reuse-chain", "提交包、发布、复用和后续种子"],
    ["P10", "problem-resolution-expert-chain", "阻碍诊断、专家组、隔离调整与提升"],
]

metrics_rows = [
    ["公开目录总源数", str(j1030_analysis["rows_master"])],
    ["带光谱性质的源", str(j1030_analysis["rows_with_spectral_properties"])],
    ["光谱或高可信红移", str(j1030_analysis["spectroscopic_or_secure_redshift"])],
    ["AGN-like 源", str(j1030_analysis["agn_like"])],
    ["obscured 候选", str(j1030_analysis["obscured_logNH_ge_22"])],
    ["heavily obscured 候选", str(j1030_analysis["heavily_obscured_logNH_ge_23"])],
    ["Compton-thick 量级候选", str(j1030_analysis["ct_level_logNH_ge_24"])],
    ["XID361 状态", "存在；低 X-ray luminosity NL-AGN 基准源，但目录 NH 不支持直接列为 heavily obscured"],
    ["P8 阀门", f"{j1030_deep['gate']['decision']}; quality_score={j1030_deep['gate']['quality_score']}"],
]

mindmap_mermaid = """```mermaid
mindmap
  root((Codex Research Loop))
    输入转化
      低结构自然语言
      normalize 补全上下文
      route 生成多链路任务图
    控制平面
      .research-loop 状态目录
      material passport
      evidence ledger
      decision log
    多通路链路
      P1 问题定义
      P2 文献证据
      P5 数据代码
      P6 分析图表
      P7 写作引用
      P8 评审完整性
      P9 提交复用
      P10 问题诊断
    共享节点
      状态读取
      深度分配
      预检分发
      物料登记
      质量阀门
      失败回流
    深循环执行
      deep-loop 生成决策
      auto-loop 消费决策
      route_next
      retry_same_route
      escalate_problem_loop
    存储与素材
      adaptive 兼容旧项目
      canonical 初始化新项目
      article chain
      data package
    证据与质量
      claim-evidence
      strict verification
      checkpoint
      handoff
    成果示例
      J1030 正式报告
      243 行公开目录
      XID235 与 XID19
      P8 quality 0.96
```"""


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        test = current + char
        bbox = draw.textbbox((0, 0), test, font=font)
        if bbox[2] - bbox[0] <= max_width or not current:
            current = test
        else:
            lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_node(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    title: str,
    lines: list[str],
    fill: str,
    outline: str,
    title_fill: str = "#0B2545",
) -> None:
    x0, y0, x1, y1 = xy
    draw.rounded_rectangle(xy, radius=20, fill=fill, outline=outline, width=3)
    title_font = get_font(27, bold=True)
    body_font = get_font(22)
    tx = x0 + 22
    ty = y0 + 16
    draw.text((tx, ty), title, font=title_font, fill=title_fill)
    y = ty + 42
    for line in lines:
        wrapped = wrap_text(draw, line, body_font, x1 - x0 - 52)
        for idx, part in enumerate(wrapped):
            prefix = "- " if idx == 0 else "  "
            draw.text((tx, y), prefix + part, font=body_font, fill="#1F2933")
            y += 30


def build_mindmap_image() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 1080
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)

    title_font = get_font(42, bold=True)
    small_font = get_font(22)
    draw.text((90, 50), "Codex Research Loop 科研工作流思维导图", font=title_font, fill="#0B2545")
    draw.text(
        (92, 110),
        "从自然语言入口到多链路科研执行、深循环阀门、问题诊断与成果交付",
        font=small_font,
        fill="#52616B",
    )

    center = (705, 415, 1095, 665)
    draw.rounded_rectangle(center, radius=34, fill="#0B2545", outline="#0B2545", width=4)
    center_title = get_font(34, bold=True)
    center_body = get_font(25)
    draw.text((756, 470), "Codex Research", font=center_title, fill="#FFFFFF")
    draw.text((820, 512), "Loop", font=center_title, fill="#FFFFFF")
    draw.text((760, 570), "项目控制平面 + 多链路路由", font=center_body, fill="#D9EAF7")

    nodes = [
        ((90, 180, 470, 340), "输入转化", ["自然语言补全", "normalize 与 route", "链路深度分配"], "#EEF6FF", "#2E74B5"),
        ((90, 390, 470, 550), "控制平面", [".research-loop", "材料护照与证据账本", "决策日志与检查点"], "#F4F7FB", "#52616B"),
        ((90, 610, 470, 820), "存储与素材", ["adaptive 兼容旧项目", "canonical 自启项目", "文章链路与数据包"], "#F4FFF7", "#2D7D46"),
        ((90, 860, 470, 1030), "问题诊断链", ["专家组评估", "scratch lab 隔离", "promotion gate 后改核心"], "#FFF7ED", "#B45309"),
        ((1330, 170, 1710, 360), "P1-P10 多通路", ["文献、数据、分析", "写作、评审、提交", "问题链 P10 回流"], "#EEF6FF", "#2E74B5"),
        ((1330, 405, 1710, 585), "深循环执行", ["deep-loop 产出决策", "auto-loop 消费决策", "Codex CLI 自启动"], "#F5F3FF", "#6D5BD0"),
        ((1330, 630, 1710, 810), "证据与质量", ["claim-evidence", "strict verification", "checkpoint/handoff"], "#F4F7FB", "#52616B"),
        ((1330, 855, 1710, 1030), "成果展示", ["J1030 正式报告", "243 行公开目录", "P8 quality=0.96"], "#FDF2F8", "#BE185D"),
    ]

    for box, title, lines, fill, outline in nodes:
        x0, y0, x1, y1 = box
        cx = center[0] if x1 < center[0] else center[2]
        cy = (center[1] + center[3]) // 2
        tx = x1 if x1 < center[0] else x0
        ty = (y0 + y1) // 2
        draw.line((cx, cy, tx, ty), fill="#9AA8B5", width=5)
        draw.ellipse((tx - 8, ty - 8, tx + 8, ty + 8), fill=outline)
        draw_node(draw, box, title, lines, fill, outline)

    img.save(MINDMAP_PNG)


md = f"""# Codex Research Loop 设计报告与 J1030 课题成果展示

报告日期：{REPORT_DATE}

插件名称：`{plugin_manifest["name"]}`
插件版本：`{plugin_manifest["version"]}`
GitHub 仓库：`https://github.com/sum331/codex-research-loop`
成果课题：`{j1030_passport["project_title"]}`

## 摘要

本报告总结 `Codex Research Loop` 的设计目标、架构、核心链路、深循环机制、无人值守执行方式和跨项目复用方法，并以 J1030 深场公开目录课题作为一次完整成果展示。该 loop 的定位不是单一自动化脚本，而是科研项目的控制平面：它把自然语言输入、项目状态、素材摄取、证据登记、链路路由、质量阀门、问题诊断和交付物管理统一进 `.research-loop/`。

在 J1030 课题中，loop 完成了从粗糙课题输入到正式科研报告的闭环：建立项目护照和存储规范，导入公开文献与目录材料，生成候选体分析表、图件和正式 DOCX 报告，并通过 P8 评审阀门。最终产物证明该系统能够支持真实科研任务，而不仅是流程设计。

## 1. 设计目标

`Codex Research Loop` 面向的是跨线程、跨项目、可长期运行的科研工作。它解决四类问题：

1. 用户输入通常是自然语言、低结构、缺少上下文，系统需要先补全为可执行工作提示。
2. 科研任务不是单链路，而是由文献、数据、方法、分析、写作、评审、发布、问题诊断等多条链路组成。
3. 长时间无人值守执行必须依赖明确阀门，而不是让代理无边界扩展。
4. 失败诊断和调整必须隔离在核心链路外，只有通过提升阀门后才能改动核心文件。

因此，loop 的设计原则是：项目状态本地化、材料和证据可追踪、路由多通路、深度可调、失败可回流、交付可验证。

## 2. 总体架构

### 2.1 思维导图

下图概括了 loop 的核心结构：自然语言输入先进入 normalize 与 route，随后由共享节点分发给 P1-P10 多条科研子链路；deep-loop 负责评估当前链路是否通过，auto-loop 负责消费该决策并继续无人值守执行；当项目出现阻碍时，P10 问题诊断链在隔离环境中生成专家评估和调整方案。

![Codex Research Loop 思维导图](assets/loop_research_mindmap.png)

{mindmap_mermaid}

### 2.2 控制平面

每个项目都有独立的 `.research-loop/` 目录，记录项目状态、材料护照、证据账本、决策日志、运行日志、深循环报告、问题诊断案例、检查点和交接文件。它不直接替代研究材料本身，而是保存控制信息和审计信息。

### 2.3 插件组成

{md_table(["模块", "实现入口", "作用"], capability_rows)}

### 2.4 多链路科研流程

{md_table(["链路", "名称", "职责"], subchain_rows)}

### 2.5 共享节点

系统在 P1-P10 子链路之外设置共享节点 N0-N9，包括状态读取、任务接收、深度分配、预检、分发、物料登记、账本更新、质量阀门、检查点交接和失败回流。这样可以避免把所有任务压进一条线性链路，也可以在失败后按问题类型回到合适子链路。

```mermaid
flowchart LR
  N0["N0 state-read"] --> N1["N1 task-intake"]
  N1 --> N2["N2 depth-assignment"]
  N2 --> N3["N3 preflight-gate"]
  N3 --> N4["N4 dispatcher"]
  N4 --> P2["P2 literature"]
  N4 --> P5["P5 data-code"]
  N4 --> P6["P6 analysis-figure"]
  N4 --> P7["P7 writing"]
  P2 --> N5["N5 artifact-registry"]
  P5 --> N5
  P6 --> N5
  P7 --> N5
  N5 --> N6["N6 ledger-update"]
  N6 --> N7["N7 quality-gate"]
  N7 -->|pass| N8["N8 checkpoint-handoff"]
  N7 -->|fail| N9["N9 feedback-dispatcher"]
  N9 --> P10["P10 problem-loop"]
```

## 3. 深循环与无人值守机制

深循环由两层组成：

1. `deep-loop`：评估当前子链路是否通过，并输出 `route_next`、`retry_same_route`、`escalate_problem_loop` 或 `pause_for_human`。
2. `auto-loop --auto-route-next`：消费 `deep-loop` 决策，并真正启动下一子链路或同链路重试。

这一区分很关键。单独运行 `deep-loop` 只会生成路由指令，不会自动执行下一链路；无人值守运行必须由 `auto-loop` 接管。推荐命令形态为：

```powershell
python C:\\Users\\<User>\\plugins\\codex-research-loop\\scripts\\research_loop.py --cwd "D:\\Project" auto-loop `
  --goal "finish current research stage" `
  --current-subchain P7 `
  --next-subchain P8 `
  --auto-route-next `
  --route-depth-budget 3 `
  --route-agent codex `
  --max-rounds 8
```

深循环不是越深越好，而是由任务难度、证据严格度、目标交付物和失败信号共同决定。课程报告类任务适合有限深循环：P5 数据复现、P6 分析图表、P7 写作引用、P8 评审完整性、P9 提交包收尾。真正需要长时间无人值守时，再开启更高的 `route-depth-budget`。

## 4. 问题诊断链设计

当项目出现阻碍时，loop 不直接修改核心文件，而是进入 P10：

1. 读取项目背景、状态、材料、路由图和现有日志。
2. 生成固定角色专家和现场构建专家，例如项目上下文分析师、测试验证工程师、隔离存储守护者、实现门禁评审员。
3. 在 `scratch/research-loop-labs/<case-id>/` 下保存诊断、测试和调整方案。
4. 通过 `problem-promote` 的提升阀门后，才允许把调整方案推广到主链路。

这一设计防止自动诊断污染主项目，尤其适合长周期科研项目和无人值守执行。

## 5. 安装与跨电脑复用

另一台电脑使用时，只需要克隆公开仓库，然后在目标项目中初始化：

```powershell
git clone https://github.com/sum331/codex-research-loop.git C:\\Users\\<User>\\plugins\\codex-research-loop
python C:\\Users\\<User>\\plugins\\codex-research-loop\\scripts\\research_loop.py --cwd "D:\\Research\\MyProject" init
python C:\\Users\\<User>\\plugins\\codex-research-loop\\scripts\\research_loop.py --cwd "D:\\Research\\MyProject" storage --style canonical --init-dirs --write
```

如果要迁移旧项目，应复制项目目录本身，并保留 `.research-loop/`。GitHub 仓库存放插件源码，不携带某个具体项目的私有材料和运行账本。

## 6. J1030 课题应用成果

### 6.1 课题定义

课题名称：`{j1030_passport["project_title"]}`
研究领域：`{j1030_passport["domain"]}`
目标场景：`{j1030_passport["profile"]["target_venue"]}`
研究问题：`{j1030_passport["research_question"]}`

### 6.2 执行链路

该课题经历了以下关键阶段：

1. P1/P2：根据课程大纲和课题草案，明确公开目录 AGN 候选体筛选问题，并登记文献元数据。
2. P5：构建 J1030 公开目录数据包、处理有效 VizieR 表格、隔离无效 bot-check 下载。
3. P6：生成 243 行 master candidate table、候选排序表和正式图件。
4. P7：形成正式科研报告 Markdown 与 DOCX。
5. P8：通过正式评审阀门，质量分数 0.96，决策为 `route_next`。

### 6.3 关键量化结果

{md_table(["指标", "结果"], metrics_rows)}

### 6.4 高优先级候选体展示

{md_table(["候选体", "分数", "类别", "红移", "光谱类型", "log NH", "log L2-10"], top5_rows)}

### 6.5 产物清单

主要产物包括：

- 正式报告 DOCX：`{case_ref(j1030_build["docx"])}`
- 正式报告 Markdown：`{case_ref(j1030_build["markdown"])}`
- Word 导出 QA PDF：`{case_ref(PROJECT_ROOT / "reports" / "j1030_formal_research_report_word_export.pdf")}`
- master candidate table：`{case_ref(j1030_analysis["outputs"]["master_table"])}`
- top20 candidate table：`{case_ref(j1030_analysis["outputs"]["top20_table"])}`
- all ranked candidate table：`{case_ref(j1030_analysis["outputs"]["all_ranked_table"])}`
- 正式图件：`{case_ref(j1030_build["figures"]["fig1"])}`, `{case_ref(j1030_build["figures"]["fig2"])}`, `{case_ref(j1030_build["figures"]["fig3"])}`

### 6.6 科研结论展示

J1030 课题证明：公开目录可以构建有物理意义的重度遮蔽 AGN 候选池。XID235 与 XID19 是最稳健的高优先级候选，XID189 是 Compton-thick 量级 NH 候选。XID361 被恢复为 X-ray-underluminous NL-AGN 基准源，但仅凭目录 NH 不能将其列为重度遮蔽确认源，因此更适合作为多波段不匹配诊断的基准对象。

## 7. 设计复盘

这轮真实课题运行暴露出一个重要工程边界：`deep-loop` 与 `auto-loop` 必须明确分层。`deep-loop` 负责评估和生成下一链路指令；`auto-loop --auto-route-next` 才负责消费该指令并继续执行。如果用户期望无人值守深循环，入口命令必须包含自动路由执行器。

这一问题反而完善了系统设计：之后应在项目 profile 中加入运行模式，例如 `delivery_mode`、`unattended_mode` 和 `terminal_stage_policy`。当任务是交付型报告时，P8/P9 通过后可以停止；当任务是探索型科研时，可以继续进入更深链路。

## 8. 后续增强方向

1. `repository-publisher`：支持 Zenodo/OSF 发布、DOI 回填和 release provenance。
2. `research-kg-builder`：把论文、主张、证据、数据、运行、图件和决策构成知识图谱。
3. `experiment-runner-plus`：记录环境哈希、资源消耗、批处理任务和失败重试策略。
4. `literature-monitor`：定期监控新论文、引用和竞争结果。
5. `semantic claim verifier`：在结构性证据检查之外，加入语义级证据支持判断。

## 9. 结论

`Codex Research Loop` 已经从一个局部 loop 升级为可跨项目复用的科研工作流插件。它的核心价值不是替代专业研究技能，而是把专业技能组织进一个可追踪、可验证、可恢复、可无人值守推进的控制系统。J1030 课题展示了该设计在真实科研课程项目中的可用性：从自然语言任务到数据处理、图表、正式报告和评审阀门，形成了完整闭环。

"""


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.append(grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_style_font(style, name: str, size: int, color: str | None = None, bold: bool | None = None) -> None:
    font = style.font
    font.name = name
    font.size = Pt(size)
    if color:
        font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_run_font(run, name: str = "Calibri", east_asia: str = "Microsoft YaHei") -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def add_paragraph(doc: Document, text: str = "", style: str | None = None):
    p = doc.add_paragraph(text, style=style)
    for run in p.runs:
        set_run_font(run)
    return p


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for idx, head in enumerate(headers):
        hdr[idx].text = head
        set_cell_shading(hdr[idx], "E8EEF5")
        set_cell_margins(hdr[idx])
        hdr[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for p in hdr[idx].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                set_run_font(run)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
            set_cell_margins(cells[idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
                    set_run_font(run)
    set_table_width(table, widths)
    add_paragraph(doc)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        run = p.add_run(item)
        set_run_font(run)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    set_style_font(doc.styles["Normal"], "Calibri", 11)
    set_style_font(doc.styles["Heading 1"], "Calibri", 16, "2E74B5", True)
    set_style_font(doc.styles["Heading 2"], "Calibri", 13, "2E74B5", True)
    set_style_font(doc.styles["Heading 3"], "Calibri", 12, "1F4D78", True)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.1

    title = add_paragraph(doc)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Codex Research Loop 设计报告与 J1030 课题成果展示")
    set_run_font(run)
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = add_paragraph(doc)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"插件版本 {plugin_manifest['version']} | 报告日期 {REPORT_DATE}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("555555")
    set_run_font(run)

    add_paragraph(doc, "摘要", style="Heading 1")
    add_paragraph(
        doc,
        "本报告总结 Codex Research Loop 的设计目标、架构、核心链路、深循环机制、无人值守执行方式和跨项目复用方法，并以 J1030 深场公开目录课题作为一次完整成果展示。该 loop 的定位不是单一自动化脚本，而是科研项目的控制平面。",
    )
    add_paragraph(
        doc,
        "在 J1030 课题中，loop 完成了从粗糙课题输入到正式科研报告的闭环：建立项目护照和存储规范，导入公开文献与目录材料，生成候选体分析表、图件和正式 DOCX 报告，并通过 P8 评审阀门。",
    )

    add_paragraph(doc, "1. 设计目标", style="Heading 1")
    add_bullets(
        doc,
        [
            "把低结构自然语言输入转化为项目化、可执行、可验证的工作提示。",
            "把文献、数据、方法、分析、写作、评审、发布和问题诊断组织为多通路链路。",
            "用通用阀门控制继续、重试、升级问题链或暂停给人类。",
            "将失败诊断隔离在核心链路外，过提升阀门后才允许改动主项目。",
        ],
    )

    add_paragraph(doc, "2. 总体架构", style="Heading 1")
    add_paragraph(doc, "2.1 思维导图", style="Heading 2")
    add_paragraph(
        doc,
        "思维导图概括了 loop 的核心结构：自然语言输入先进入 normalize 与 route，随后由共享节点分发给 P1-P10 多条科研子链路；deep-loop 负责评估当前链路是否通过，auto-loop 负责消费该决策并继续无人值守执行；当项目出现阻碍时，P10 问题诊断链在隔离环境中生成专家评估和调整方案。",
    )
    if MINDMAP_PNG.exists():
        doc.add_picture(str(MINDMAP_PNG), width=Inches(6.3))
        cap = add_paragraph(doc, "图 A  Codex Research Loop 科研工作流思维导图")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cap.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string("555555")

    add_paragraph(doc, "2.2 插件组成", style="Heading 2")
    add_table(doc, ["模块", "实现入口", "作用"], capability_rows, [1900, 2500, 4960])

    add_paragraph(doc, "2.3 多链路科研流程", style="Heading 2")
    add_table(doc, ["链路", "名称", "职责"], subchain_rows, [900, 3100, 5360])

    add_paragraph(doc, "2.4 共享节点", style="Heading 2")
    add_paragraph(
        doc,
        "P1-P10 子链路之外，系统设置 N0-N9 共享节点：状态读取、任务接收、深度分配、预检、分发、物料登记、账本更新、质量阀门、检查点交接和失败回流。共享节点让多任务可以共用入口、证据和阀门，而不是被压成一条线性流程。",
    )

    add_paragraph(doc, "3. 深循环与无人值守机制", style="Heading 1")
    add_paragraph(
        doc,
        "深循环由两层组成：deep-loop 负责评估当前子链路并生成 route_next、retry_same_route、escalate_problem_loop 或 pause_for_human；auto-loop --auto-route-next 负责消费 deep-loop 决策并真正启动下一子链路或同链路重试。",
    )
    add_paragraph(
        doc,
        "这一分层是关键工程边界。单独运行 deep-loop 只会生成路由指令，不会自动进入下一链路。无人值守运行必须由 auto-loop 接管，并配置 route-agent codex 或自定义 route-agent-command。",
    )

    add_paragraph(doc, "4. 问题诊断链设计", style="Heading 1")
    add_bullets(
        doc,
        [
            "读取项目背景、状态、材料、路由图和现有日志。",
            "生成固定角色专家和现场构建专家。",
            "在 scratch/research-loop-labs/<case-id>/ 中保存诊断、测试和调整方案。",
            "通过 problem-promote 提升阀门后，才允许把调整方案推广到主链路。",
        ],
    )

    add_paragraph(doc, "5. 跨电脑复用", style="Heading 1")
    add_paragraph(
        doc,
        "另一台电脑只需克隆公开仓库，并在目标项目中运行 init 与 storage。GitHub 仓库存放插件源码，不携带某个项目的私有材料和运行账本；迁移旧项目时需要连同 .research-loop/ 一起复制项目目录。",
    )

    doc.add_page_break()
    add_paragraph(doc, "6. J1030 课题应用成果", style="Heading 1")
    add_paragraph(doc, "6.1 课题定义", style="Heading 2")
    add_table(
        doc,
        ["字段", "内容"],
        [
            ["课题名称", j1030_passport["project_title"]],
            ["研究领域", j1030_passport["domain"]],
            ["目标场景", j1030_passport["profile"]["target_venue"]],
            ["研究问题", j1030_passport["research_question"]],
        ],
        [1800, 7560],
    )

    add_paragraph(doc, "6.2 执行链路", style="Heading 2")
    add_bullets(
        doc,
        [
            "P1/P2：根据课程大纲和课题草案明确公开目录 AGN 候选体筛选问题，并登记文献元数据。",
            "P5：构建 J1030 公开目录数据包、处理有效 VizieR 表格、隔离无效 bot-check 下载。",
            "P6：生成 243 行 master candidate table、候选排序表和正式图件。",
            "P7：形成正式科研报告 Markdown 与 DOCX。",
            "P8：通过正式评审阀门，质量分数 0.96，决策为 route_next。",
        ],
    )

    doc.add_page_break()
    add_paragraph(doc, "6.3 关键量化结果", style="Heading 2")
    add_table(doc, ["指标", "结果"], metrics_rows, [3000, 6360])

    add_paragraph(doc, "6.4 高优先级候选体", style="Heading 2")
    add_table(doc, ["候选体", "分数", "类别", "红移", "光谱类型", "log NH", "log L2-10"], top5_rows, [900, 700, 1900, 850, 1200, 950, 1100])

    add_paragraph(doc, "6.5 成果图件", style="Heading 2")
    for label, fig_path in [
        ("图 1  红移与 X-ray luminosity 候选分布", Path(j1030_build["figures"]["fig1"])),
        ("图 2  红移与 NH 候选分布", Path(j1030_build["figures"]["fig2"])),
        ("图 3  NH 分布", Path(j1030_build["figures"]["fig3"])),
    ]:
        if fig_path.exists():
            doc.add_picture(str(fig_path), width=Inches(6.1))
            cap = add_paragraph(doc, label)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor.from_string("555555")

    add_paragraph(doc, "6.6 主要交付物", style="Heading 2")
    add_bullets(
        doc,
        [
            f"正式报告 DOCX：{case_ref(j1030_build['docx'])}",
            f"正式报告 Markdown：{case_ref(j1030_build['markdown'])}",
            f"Word 导出 QA PDF：{case_ref(PROJECT_ROOT / 'reports' / 'j1030_formal_research_report_word_export.pdf')}",
            f"master candidate table：{case_ref(j1030_analysis['outputs']['master_table'])}",
            f"top20 candidate table：{case_ref(j1030_analysis['outputs']['top20_table'])}",
            f"all ranked candidate table：{case_ref(j1030_analysis['outputs']['all_ranked_table'])}",
        ],
    )

    add_paragraph(doc, "6.7 科研结论展示", style="Heading 2")
    add_paragraph(
        doc,
        "J1030 课题证明：公开目录可以构建有物理意义的重度遮蔽 AGN 候选池。XID235 与 XID19 是最稳健的高优先级候选，XID189 是 Compton-thick 量级 NH 候选。XID361 被恢复为 X-ray-underluminous NL-AGN 基准源，但仅凭目录 NH 不能将其列为重度遮蔽确认源，因此更适合作为多波段不匹配诊断的基准对象。",
    )

    add_paragraph(doc, "7. 设计复盘与后续增强", style="Heading 1")
    add_paragraph(
        doc,
        "这轮真实课题运行暴露出一个重要工程边界：deep-loop 与 auto-loop 必须明确分层。deep-loop 负责评估和生成下一链路指令；auto-loop --auto-route-next 才负责消费该指令并继续执行。",
    )
    add_bullets(
        doc,
        [
            "repository-publisher：支持 Zenodo/OSF 发布、DOI 回填和 release provenance。",
            "research-kg-builder：把论文、主张、证据、数据、运行、图件和决策构成知识图谱。",
            "experiment-runner-plus：记录环境哈希、资源消耗、批处理任务和失败重试策略。",
            "literature-monitor：定期监控新论文、引用和竞争结果。",
            "semantic claim verifier：在结构性证据检查之外，加入语义级证据支持判断。",
        ],
    )

    add_paragraph(doc, "8. 结论", style="Heading 1")
    add_paragraph(
        doc,
        "Codex Research Loop 已经从一个局部 loop 升级为可跨项目复用的科研工作流插件。它的核心价值不是替代专业研究技能，而是把专业技能组织进一个可追踪、可验证、可恢复、可无人值守推进的控制系统。J1030 课题展示了该设计在真实科研课程项目中的可用性。",
    )

    doc.add_section(WD_SECTION.CONTINUOUS)
    doc.save(OUT_DOCX)
    shutil.copyfile(OUT_DOCX, OUT_DOCX_LEGACY)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_mindmap_image()
    OUT_MD.write_text(md.rstrip() + "\n", encoding="utf-8")
    build_docx()
    summary = {
        "date": REPORT_DATE,
        "markdown": repo_ref(OUT_MD),
        "docx": repo_ref(OUT_DOCX),
        "legacy_docx": repo_ref(OUT_DOCX_LEGACY),
        "mindmap": repo_ref(MINDMAP_PNG),
        "plugin_version": plugin_manifest["version"],
        "case_project": "$env:J1030_PROJECT_ROOT",
        "j1030_rows": j1030_analysis["rows_master"],
        "j1030_gate": j1030_deep["gate"],
    }
    (OUT_DIR / "codex_research_loop_design_and_j1030_case_report_build_summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
